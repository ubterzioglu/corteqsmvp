#!/usr/bin/env python3
"""
B4 · role_structure tablosunu Excel'den seed et.

Kullanım:
  python scripts/seed-role-structure-from-excel.py

Bağımlılıklar:
  pip install openpyxl

Çıktı:
  supabase/migrations/applied/20261007130000_role_structure_seed.sql
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pathlib import Path
from openpyxl import load_workbook

# Excel'den oku
excel_path = Path(__file__).parent.parent / "CorteQS_Rol_Tablosu_v2.xlsx"
output_path = Path(__file__).parent.parent / "supabase/migrations/applied/20261007130000_role_structure_seed.sql"

print(f"📖 Excel okunuyor: {excel_path}")
wb = load_workbook(excel_path, data_only=True)
ws = wb["Yeni Rol Yapısı"]

rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[1]:  # Ana Rol boşsa atla
        continue
    
    ana_rol = str(row[1] or "").strip()
    alt_rol = str(row[3] or "").strip()
    uzmanlik = str(row[5] or "").strip()
    yeni_kod = str(row[6] or "").strip()
    eski_kod = str(row[7] or "").strip() if row[7] and str(row[7]).strip() not in ("—", "-", "") else None
    durum = "onaylandi" if str(row[9] or "").strip().lower() == "onaylandı" else "oneri"
    
    if not yeni_kod:
        continue
    
    rows.append({
        "ana_rol": ana_rol,
        "alt_rol": alt_rol,
        "uzmanlik": uzmanlik,
        "yeni_kod": yeni_kod,
        "eski_kod": eski_kod,
        "durum": durum,
    })

print(f"✅ {len(rows)} satır okundu")

# SQL üret
def escape_sql(s: str) -> str:
    """SQL string'de kaçış."""
    return s.replace("'", "''")

sql_lines = [
    "-- B4 · role_structure seed (Excel'den üretildi)",
    "-- Bu dosya otomatik üretildi: python scripts/seed-role-structure-from-excel.py",
    "",
    "begin;",
    "",
]

for r in rows:
    eski_kod_sql = f"'{escape_sql(r['eski_kod'])}'" if r['eski_kod'] else "null"
    sql_lines.append(
        f"insert into public.role_structure (ana_rol, alt_rol, uzmanlik, yeni_kod, eski_kod, durum) "
        f"values ('{escape_sql(r['ana_rol'])}', '{escape_sql(r['alt_rol'])}', '{escape_sql(r['uzmanlik'])}', "
        f"'{escape_sql(r['yeni_kod'])}', {eski_kod_sql}, '{r['durum']}') "
        f"on conflict (yeni_kod) do nothing;"
    )

sql_lines.extend([
    "",
    "commit;",
])

# Yaz
output_path.write_text("\n".join(sql_lines), encoding="utf-8")
print(f"✅ Seed SQL yazıldı: {output_path}")
print(f"   Toplam: {len(rows)} satır")
print(f"   Ana rol: {len(set(r['ana_rol'] for r in rows))}")
print(f"   Alt rol: {len(set(r['alt_rol'] for r in rows))}")
print(f"   ÖNERİ: {sum(1 for r in rows if r['durum'] == 'oneri')}")
